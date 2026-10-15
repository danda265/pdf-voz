"""Modo OCR: texto reconhecido em PDFs que sao so imagem vira caixas de texto editaveis."""
from __future__ import annotations

import io
import os
import unicodedata

import fitz
import pytest
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import pdf_to_pptx
import pdf_to_pptx_ocr as ocr_mod
from pdf_to_pptx import MODO_EDITAVEL, MODO_OCR, ConversaoError, converter_pdf_para_pptx
from pdf_to_pptx_ocr import Linha, _junta_em_paragrafos, _sem_marcadores, reconhecer_e_limpar

ARIAL = r'C:\Windows\Fonts\arial.ttf'
ARIAL_NEGRITO = r'C:\Windows\Fonts\arialbd.ttf'

precisa_de_ocr = pytest.mark.skipif(
    not (ocr_mod.tesseract_disponivel() and os.path.exists(ARIAL) and os.path.exists(ARIAL_NEGRITO)),
    reason='Tesseract embutido ou fontes Arial ausentes')


def _sem_acento(texto: str) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn').lower()


def _slide(linhas, fundo=(250, 250, 255), tamanho=(1600, 900)):
    """Imagem de fundo liso com texto: [(x, y, texto, tamanho_px, cor, negrito), ...]."""
    imagem = Image.new('RGB', tamanho, fundo)
    desenho = ImageDraw.Draw(imagem)
    for x, y, texto, px, cor, negrito in linhas:
        desenho.text((x, y), texto, font=ImageFont.truetype(ARIAL_NEGRITO if negrito else ARIAL, px), fill=cor)
    return imagem


def _pdf_da_imagem(tmp_path, imagem, nome='slide.pdf'):
    memoria = io.BytesIO()
    imagem.save(memoria, 'PNG')
    origem = tmp_path / nome
    doc = fitz.open()
    pagina = doc.new_page(width=imagem.width / 2, height=imagem.height / 2)
    pagina.insert_image(pagina.rect, stream=memoria.getvalue())
    doc.save(str(origem))
    doc.close()
    return origem


def _linha(y, tamanho, x0=100, largura=500, altura=30, texto='texto de teste'):
    return Linha(texto=texto, x0=x0, y0=y, x1=x0 + largura, y1=y + altura, maior_palavra=altura,
                 confianca=95.0, tamanho_px=tamanho)


# ── sem depender do OCR ──────────────────────────────────────────────────────

def test_linhas_de_mesmo_tamanho_uma_sob_a_outra_viram_um_paragrafo():
    paragrafos = _junta_em_paragrafos([_linha(100, 30), _linha(140, 30), _linha(180, 30)])
    assert len(paragrafos) == 1 and len(paragrafos[0].linhas) == 3


def test_tamanhos_diferentes_ou_longe_ou_desalinhados_ficam_separados():
    grande_e_pequena = _junta_em_paragrafos([_linha(100, 42), _linha(145, 30)])
    assert len(grande_e_pequena) == 2
    longe = _junta_em_paragrafos([_linha(100, 30), _linha(400, 30)])
    assert len(longe) == 2
    desalinhadas = _junta_em_paragrafos([_linha(100, 30), _linha(140, 30, x0=700, largura=100)])
    assert len(desalinhadas) == 2


def test_linhas_centralizadas_sao_detectadas_como_centralizadas():
    p = _junta_em_paragrafos([_linha(100, 30, x0=300, largura=400), _linha(140, 30, x0=350, largura=300)])[0]
    ocr_mod._ajusta_paragrafo(p)
    assert p.centralizado


def test_icone_no_comeco_da_linha_nao_vira_texto():
    # (texto, esquerda, topo, largura, altura, confianca): o 'Q' e o icone de um monitor
    palavras = [('Q', 10, 10, 20, 30, 84.0), ('Introdução', 80, 10, 200, 30, 95.0), ('às', 288, 10, 40, 30, 95.0)]
    assert [p[0] for p in _sem_marcadores(palavras)] == ['Introdução', 'às']


def test_palavra_curta_colada_no_texto_e_simbolo_isolado():
    colada = [('A', 10, 10, 20, 30, 93.0), ('estrutura', 40, 10, 180, 30, 95.0)]    # vao de 10 px: e uma palavra
    assert [p[0] for p in _sem_marcadores(colada)] == ['A', 'estrutura']
    simbolo = [('*', 10, 10, 12, 30, 52.0), ('Cultura', 24, 10, 150, 30, 95.0)]       # simbolo sai ate colado
    assert [p[0] for p in _sem_marcadores(simbolo)] == ['Cultura']
    sobra_uma = [('Coluna', 10, 10, 120, 30, 95.0)]
    assert [p[0] for p in _sem_marcadores(sobra_uma)] == ['Coluna']


def test_so_apaga_o_texto_se_o_fundo_for_liso_e_so_a_tinta():
    imagem = Image.new('RGB', (400, 200), (240, 240, 250))
    ImageDraw.Draw(imagem).text((50, 80), 'texto', font=ImageFont.load_default(size=40), fill=(0, 0, 0))
    antes = imagem.copy()
    assert ocr_mod._apaga_texto(imagem, (50, 80, 150, 120), 6, 12) is True
    assert min(imagem.crop((40, 70, 170, 140)).tobytes()) >= 235          # nenhuma tinta escura sobrou
    assert imagem.getpixel((5, 5)) == antes.getpixel((5, 5))               # fora da area, nada mudou


def test_fundo_com_listras_nao_e_apagado():
    imagem = Image.new('RGB', (400, 200), (255, 255, 255))
    desenho = ImageDraw.Draw(imagem)
    for x in range(0, 400, 12):
        desenho.rectangle((x, 0, x + 5, 200), fill=(30, 60, 200))
    antes = imagem.tobytes()
    assert ocr_mod._apaga_texto(imagem, (100, 80, 250, 120), 6, 12) is False
    assert imagem.tobytes() == antes                                       # a imagem so e alterada se o fundo for liso


def test_area_que_e_mais_tinta_que_fundo_nao_e_texto():
    imagem = Image.new('RGB', (400, 200), (255, 255, 255))
    ImageDraw.Draw(imagem).rectangle((90, 70, 260, 130), fill=(0, 0, 0))   # um bloco preto, nao letras
    antes = imagem.tobytes()
    assert ocr_mod._apaga_texto(imagem, (100, 80, 250, 120), 6, 12) is False
    assert imagem.tobytes() == antes


# ── com o OCR de verdade ─────────────────────────────────────────────────────

@precisa_de_ocr
def test_le_o_texto_com_acentos_e_apaga_da_imagem():
    imagem = _slide([(100, 100, 'Linhas, colunas e células', 64, (20, 40, 90), False),
                     (100, 300, 'Organizadas horizontalmente às quartas', 40, (60, 70, 90), False)])
    resultado = reconhecer_e_limpar(imagem)
    textos = [p.texto for p in resultado.paragrafos]
    assert any('células' in t for t in textos), textos           # com so portugues os acentos sobrevivem
    assert any('às' in t for t in textos), textos
    assert min(resultado.imagem.crop((90, 90, 900, 190)).tobytes()) >= 240   # o texto saiu da imagem
    assert [p.texto for p in resultado.paragrafos] == sorted(
        [p.texto for p in resultado.paragrafos], key=lambda t: textos.index(t))


@precisa_de_ocr
def test_guarda_a_cor_e_o_negrito_de_cada_trecho_da_linha():
    imagem = _slide([(100, 100, 'Cada coluna deve representar', 56, (30, 50, 110), False)])
    ImageDraw.Draw(imagem).text((100, 200), 'um', font=ImageFont.truetype(ARIAL, 56), fill=(30, 50, 110))
    ImageDraw.Draw(imagem).text((200, 200), 'tipo de informação.', font=ImageFont.truetype(ARIAL_NEGRITO, 56),
                                fill=(0, 90, 230))
    resultado = reconhecer_e_limpar(imagem)
    trechos = [t for p in resultado.paragrafos for t in p.trechos]
    negrito = [t for t in trechos if t.negrito]
    normal = [t for t in trechos if not t.negrito]
    assert negrito and normal
    assert any('informa' in t.texto for t in negrito)
    cor = negrito[-1].cor
    assert cor[2] > 180 and cor[0] < 80                             # o azul do trecho em negrito


@precisa_de_ocr
def test_texto_sobre_fundo_com_ruido_fica_na_imagem():
    import random
    sorteio = random.Random(7)
    imagem = Image.new('RGB', (1600, 900))
    imagem.putdata([(v, v, v) for v in (110 + sorteio.randint(-45, 45) for _ in range(1600 * 900))])
    ImageDraw.Draw(imagem).text((120, 120), 'Texto sobre foto', font=ImageFont.truetype(ARIAL_NEGRITO, 80),
                                fill=(255, 255, 255))
    antes = imagem.tobytes()
    resultado = reconhecer_e_limpar(imagem)
    assert resultado.paragrafos == []                       # o fundo nao e liso: refazer estragaria a imagem
    assert resultado.imagem.tobytes() == antes


# ── no conversor ─────────────────────────────────────────────────────────────

def _caixas_de_texto(slide):
    return [s for s in slide.shapes if s.has_text_frame and s.name.startswith('Texto reconhecido')]


@precisa_de_ocr
def test_pdf_so_de_imagem_vira_figura_sem_o_texto_mais_caixas_editaveis(tmp_path):
    imagem = _slide([(100, 100, 'Planilhas e tabelas', 80, (20, 40, 90), True),
                     (100, 320, 'Toda planilha é formada por linhas', 44, (70, 80, 100), False),
                     (100, 380, 'e colunas organizadas.', 44, (70, 80, 100), False)])
    origem = _pdf_da_imagem(tmp_path, imagem)
    destino = tmp_path / 'ocr.pptx'
    resultado = converter_pdf_para_pptx(str(origem), str(destino), MODO_OCR)
    slide = Presentation(str(destino)).slides[0]
    caixas = _caixas_de_texto(slide)
    textos = ' '.join(c.text_frame.text for c in caixas)
    assert 'planilhas' in _sem_acento(textos) and 'colunas' in _sem_acento(textos)
    assert resultado.paginas_com_ocr == 1 and resultado.paragrafos_ocr == len(caixas)
    figuras = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(figuras) == 1 and list(slide.shapes)[0].shape_type == MSO_SHAPE_TYPE.PICTURE   # a imagem fica por baixo
    fundo = Image.open(io.BytesIO(figuras[0].image.blob)).convert('RGB')
    assert min(fundo.crop((90, 90, 900, 200)).tobytes()) >= 240           # o texto sumiu da imagem


@precisa_de_ocr
def test_paragrafo_de_varias_linhas_mantem_as_quebras_do_original(tmp_path):
    imagem = _slide([(100, 100, 'Toda planilha é formada por', 48, (70, 80, 100), False),
                     (100, 165, 'linhas e colunas e células', 48, (70, 80, 100), False)])
    origem = _pdf_da_imagem(tmp_path, imagem)
    destino = tmp_path / 'ocr.pptx'
    converter_pdf_para_pptx(str(origem), str(destino), MODO_OCR)
    caixa = _caixas_de_texto(Presentation(str(destino)).slides[0])[0]
    xml = etree.tostring(caixa._element).decode()
    assert '<a:br' in xml and caixa.text_frame.word_wrap is False          # quebra manual, nao refluxo do PowerPoint
    assert len(caixa.text_frame.paragraphs) == 1                           # continua sendo um paragrafo so


@precisa_de_ocr
def test_o_modo_editavel_nao_roda_ocr(tmp_path):
    imagem = _slide([(100, 100, 'Planilhas e tabelas', 80, (20, 40, 90), True)])
    destino = tmp_path / 'ed.pptx'
    resultado = converter_pdf_para_pptx(str(_pdf_da_imagem(tmp_path, imagem)), str(destino), MODO_EDITAVEL)
    assert resultado.paginas_com_ocr == 0 and _caixas_de_texto(Presentation(str(destino)).slides[0]) == []


def test_pagina_com_texto_de_verdade_nao_passa_pelo_ocr(sample_pdf, tmp_path):
    destino = tmp_path / 'ocr.pptx'
    resultado = converter_pdf_para_pptx(str(sample_pdf), str(destino), MODO_OCR)
    assert resultado.paginas_com_ocr == 0 and resultado.caixas_de_texto > 0


@precisa_de_ocr
def test_pagina_de_imagem_sem_texto_reconhecivel_segue_como_figura(tmp_path):
    imagem = Image.new('RGB', (1600, 900), (230, 240, 250))
    ImageDraw.Draw(imagem).ellipse((300, 200, 900, 700), fill=(40, 120, 200))
    destino = tmp_path / 'ocr.pptx'
    resultado = converter_pdf_para_pptx(str(_pdf_da_imagem(tmp_path, imagem)), str(destino), MODO_OCR)
    slide = Presentation(str(destino)).slides[0]
    assert resultado.paginas_com_ocr == 0
    assert [s.shape_type for s in slide.shapes] == [MSO_SHAPE_TYPE.PICTURE]


def test_sem_tesseract_o_modo_ocr_da_mensagem_clara(sample_pdf, tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_to_pptx, 'tesseract_disponivel', lambda: False)
    with pytest.raises(ConversaoError, match='OCR'):
        converter_pdf_para_pptx(str(sample_pdf), str(tmp_path / 's.pptx'), MODO_OCR)
    assert not (tmp_path / 's.pptx').exists()


def test_erro_inesperado_no_ocr_cai_para_a_figura_sem_quebrar(tmp_path, monkeypatch):
    def quebra(imagem):
        raise RuntimeError('tesseract travou')
    monkeypatch.setattr(pdf_to_pptx, 'tesseract_disponivel', lambda: True)
    monkeypatch.setattr(pdf_to_pptx, 'reconhecer_e_limpar', quebra)
    imagem = Image.new('RGB', (800, 450), (240, 240, 250))
    destino = tmp_path / 'ocr.pptx'
    resultado = converter_pdf_para_pptx(str(_pdf_da_imagem(tmp_path, imagem)), str(destino), MODO_OCR)
    assert resultado.paginas_com_ocr == 0
    assert [s.shape_type for s in Presentation(str(destino)).slides[0].shapes] == [MSO_SHAPE_TYPE.PICTURE]
