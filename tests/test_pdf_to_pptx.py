from __future__ import annotations

import fitz
import pytest
from pptx import Presentation
from pptx.util import Emu

from pdf_to_pptx import (
    MODO_EDITAVEL, MODO_IMAGEM, ConversaoError, _limpa, converter_pdf_para_pptx,
)


def _textos(slide) -> list:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def _figuras(slide) -> list:
    return [s for s in slide.shapes if s.shape_type == 13]   # MSO_SHAPE_TYPE.PICTURE


def _pixmap_da_figura(figura) -> fitz.Pixmap:
    return fitz.Pixmap(figura.image.blob)


def test_modo_imagem_um_slide_por_pagina_so_com_figura(multipage_pdf, tmp_path):
    destino = tmp_path / 'saida.pptx'
    resultado = converter_pdf_para_pptx(str(multipage_pdf), str(destino), MODO_IMAGEM)
    apresentacao = Presentation(str(destino))
    assert resultado.paginas == 3 and len(apresentacao.slides) == 3
    for slide in apresentacao.slides:
        assert len(_figuras(slide)) == 1
        assert _textos(slide) == []
    assert resultado.caixas_de_texto == 0


def test_modo_editavel_devolve_o_texto_em_caixas(sample_pdf, tmp_path):
    destino = tmp_path / 'saida.pptx'
    resultado = converter_pdf_para_pptx(str(sample_pdf), str(destino), MODO_EDITAVEL)
    slide = Presentation(str(destino)).slides[0]
    textos = _textos(slide)
    assert 'Ola mundo' in textos
    assert 'Segunda linha de teste' in textos
    assert resultado.caixas_de_texto == len(textos)
    assert resultado.paginas_sem_texto_editavel == 0


def test_modo_editavel_tira_o_texto_do_fundo(sample_pdf, tmp_path):
    """Sem isso o texto apareceria duas vezes: na imagem e na caixa."""
    destino = tmp_path / 'saida.pptx'
    converter_pdf_para_pptx(str(sample_pdf), str(destino), MODO_EDITAVEL)
    fundo = _pixmap_da_figura(_figuras(Presentation(str(destino)).slides[0])[0])
    assert min(fundo.samples) == 255, 'o fundo ainda tem pixels escuros (texto nao foi removido)'


def test_modo_editavel_preserva_desenhos_sob_o_texto(tmp_path):
    origem = tmp_path / 'cor.pdf'
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    page.draw_rect(fitz.Rect(20, 20, 280, 120), color=None, fill=(1, 0, 0))
    page.insert_text((40, 70), 'sobre o vermelho', fontsize=14, fontname='helv', color=(1, 1, 1))
    doc.save(str(origem))
    doc.close()
    destino = tmp_path / 'saida.pptx'
    converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    slide = Presentation(str(destino)).slides[0]
    assert 'sobre o vermelho' in _textos(slide)
    fundo = _pixmap_da_figura(_figuras(slide)[0])
    x, y = int(fundo.width * 40 / 300), int(fundo.height * 70 / 200)
    r, g, b = fundo.pixel(x, y)[:3]
    assert r > 200 and g < 60 and b < 60, f'o retangulo vermelho sumiu junto com o texto: {(r, g, b)}'


def test_o_pdf_original_nao_e_alterado(sample_pdf, tmp_path):
    antes = sample_pdf.read_bytes()
    converter_pdf_para_pptx(str(sample_pdf), str(tmp_path / 'saida.pptx'), MODO_EDITAVEL)
    assert sample_pdf.read_bytes() == antes


def test_texto_invisivel_do_ocr_nao_vira_caixa(tmp_path):
    origem = tmp_path / 'escaneado.pdf'
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    page.insert_text((40, 70), 'texto da camada de ocr', fontsize=14, fontname='helv', render_mode=3)
    doc.save(str(origem))
    doc.close()
    resultado = converter_pdf_para_pptx(str(origem), str(tmp_path / 'saida.pptx'), MODO_EDITAVEL)
    assert resultado.caixas_de_texto == 0
    assert resultado.paginas_sem_texto_editavel == 1


def test_pagina_girada_fica_so_como_imagem(tmp_path):
    origem = tmp_path / 'girado.pdf'
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    page.insert_text((40, 70), 'pagina de lado', fontsize=14, fontname='helv')
    page.set_rotation(90)
    doc.save(str(origem))
    doc.close()
    destino = tmp_path / 'saida.pptx'
    resultado = converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    assert resultado.caixas_de_texto == 0 and resultado.paginas_sem_texto_editavel == 1
    assert len(_figuras(Presentation(str(destino)).slides[0])) == 1


def test_paginas_de_tamanhos_diferentes_cabem_no_slide(tmp_path):
    origem = tmp_path / 'mistas.pdf'
    doc = fitz.open()
    doc.new_page(width=600, height=400).insert_text((50, 80), 'paisagem', fontsize=14, fontname='helv')
    doc.new_page(width=300, height=500).insert_text((50, 80), 'retrato', fontsize=14, fontname='helv')
    doc.save(str(origem))
    doc.close()
    destino = tmp_path / 'saida.pptx'
    converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    apresentacao = Presentation(str(destino))
    for slide in apresentacao.slides:
        for forma in slide.shapes:
            assert forma.left >= 0 and forma.top >= 0
            assert forma.left + forma.width <= apresentacao.slide_width + Emu(2)
            assert forma.top + forma.height <= apresentacao.slide_height + Emu(2)


def test_pagina_em_branco_vira_slide_so_com_fundo(blank_page_pdf, tmp_path):
    destino = tmp_path / 'saida.pptx'
    resultado = converter_pdf_para_pptx(str(blank_page_pdf), str(destino), MODO_EDITAVEL)
    assert len(Presentation(str(destino)).slides) == 1
    assert resultado.caixas_de_texto == 0 and resultado.paginas_sem_texto_editavel == 1


def test_progresso_chamado_uma_vez_por_pagina(multipage_pdf, tmp_path):
    chamadas = []
    converter_pdf_para_pptx(str(multipage_pdf), str(tmp_path / 's.pptx'), MODO_IMAGEM,
                            lambda atual, total: chamadas.append((atual, total)))
    assert chamadas == [(1, 3), (2, 3), (3, 3)]


def test_nao_deixa_arquivo_parcial_nem_cria_saida_com_erro(tmp_path):
    origem = tmp_path / 'ruim.pdf'
    origem.write_bytes(b'isto nao e um pdf')
    destino = tmp_path / 'saida.pptx'
    with pytest.raises(ConversaoError):
        converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    assert not destino.exists()
    assert list(tmp_path.glob('*.parcial')) == []


def test_pdf_com_senha_da_mensagem_clara(sample_pdf, tmp_path):
    protegido = tmp_path / 'protegido.pdf'
    doc = fitz.open(str(sample_pdf))
    doc.save(str(protegido), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='segredo', owner_pw='dono')
    doc.close()
    with pytest.raises(ConversaoError, match='senha'):
        converter_pdf_para_pptx(str(protegido), str(tmp_path / 's.pptx'))


def test_recusa_arquivo_que_nao_e_pdf_e_destino_igual_a_origem(sample_pdf, tmp_path):
    texto = tmp_path / 'a.txt'
    texto.write_text('x')
    with pytest.raises(ConversaoError):
        converter_pdf_para_pptx(str(texto), str(tmp_path / 's.pptx'))
    with pytest.raises(ConversaoError):
        converter_pdf_para_pptx(str(sample_pdf), str(sample_pdf))
    with pytest.raises(ConversaoError):
        converter_pdf_para_pptx(str(sample_pdf), str(tmp_path / 's.pptx'), 'inventado')


def test_caracteres_de_controle_sao_removidos():
    assert _limpa('a\x00b\x01c\x1fd') == 'abcd'
    assert _limpa('linha\tcom tab') == 'linha\tcom tab'
