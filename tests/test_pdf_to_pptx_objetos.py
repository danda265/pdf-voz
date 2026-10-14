"""Formas, imagens e recortes do PDF recriados como objetos editaveis do PowerPoint."""
from __future__ import annotations

import io

import fitz
import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

import pdf_to_pptx
from pdf_to_pptx import MODO_EDITAVEL, converter_pdf_para_pptx
from pdf_to_pptx_objetos import pixels_diferentes

VERMELHO, AZUL, VERDE = (1, 0, 0), (0, 0, 1), (0, 0.6, 0)


def _converte(tmp_path, desenhar, largura=400, altura=300, nome='t'):
    origem = tmp_path / f'{nome}.pdf'
    doc = fitz.open()
    page = doc.new_page(width=largura, height=altura)
    desenhar(page)
    doc.save(str(origem))
    doc.close()
    destino = tmp_path / f'{nome}.pptx'
    resultado = converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    return resultado, Presentation(str(destino)).slides[0]


def _png(rgba) -> bytes:
    imagem = Image.new('RGBA', (60, 40), rgba)
    memoria = io.BytesIO()
    imagem.save(memoria, 'PNG')
    return memoria.getvalue()


def _tipos(slide):
    return [s.shape_type for s in slide.shapes]


def _xml(forma) -> str:
    return etree.tostring(forma._element).decode()


# ── formas ───────────────────────────────────────────────────────────────────

def test_cada_tipo_de_forma_vira_o_objeto_certo(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(20, 20, 120, 80), color=None, fill=VERMELHO)                # retangulo
        page.draw_rect(fitz.Rect(140, 20, 260, 80), color=None, fill=AZUL, radius=0.25)      # cantos redondos
        page.draw_oval(fitz.Rect(280, 20, 380, 80), color=None, fill=VERDE)                  # elipse
        page.draw_polyline([(20, 200), (80, 120), (140, 200)], color=None, fill=VERMELHO, closePath=True)
        page.draw_line((160, 250), (380, 250), color=AZUL, width=4)                          # linha
    resultado, slide = _converte(tmp_path, desenha)
    formas = {s.name.rsplit(' ', 1)[0]: s for s in slide.shapes}
    assert resultado.formas == 5 and resultado.imagens == 0
    assert formas['Retângulo'].auto_shape_type == MSO_SHAPE.RECTANGLE
    assert formas['Retângulo arredondado'].auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert formas['Elipse'].auto_shape_type == MSO_SHAPE.OVAL
    assert formas['Forma livre'].shape_type == MSO_SHAPE_TYPE.FREEFORM
    assert formas['Linha'].shape_type == MSO_SHAPE_TYPE.LINE
    assert str(formas['Retângulo'].fill.fore_color.rgb) == 'FF0000'


def test_raio_do_canto_arredondado_e_preservado(tmp_path):
    _, slide = _converte(tmp_path, lambda p: p.draw_rect(
        fitz.Rect(20, 20, 220, 120), color=None, fill=AZUL, radius=0.25))
    forma = next(s for s in slide.shapes if s.name.startswith('Retângulo arredondado'))
    assert forma.adjustments[0] == pytest.approx(0.25, abs=0.02)      # raio = 1/4 do lado menor


def test_forma_com_curvas_vira_forma_livre_editavel(tmp_path):
    def desenha(page):
        page.draw_bezier((40, 200), (80, 40), (160, 360), (200, 200), )
        page.draw_curve((220, 200), (260, 100), (320, 200), )
    resultado, slide = _converte(tmp_path, desenha)
    livres = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.FREEFORM]
    assert livres and 'cubicBezTo' in _xml(livres[0])


def test_contorno_largura_cor_e_tracejado(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(40, 40, 200, 140), color=AZUL, fill=None, width=3, dashes='[6 3] 0')
    _, slide = _converte(tmp_path, desenha)
    forma = next(s for s in slide.shapes if s.name.startswith('Retângulo'))
    assert forma.line.width == 3 * 12700 and str(forma.line.color.rgb) == '0000FF'
    assert 'custDash' in _xml(forma)
    assert forma.fill.type == 5 or 'noFill' in _xml(forma)             # sem preenchimento


def test_transparencia_do_preenchimento_e_preservada(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(40, 40, 200, 140), color=None, fill=VERMELHO, fill_opacity=0.4)
    _, slide = _converte(tmp_path, desenha)
    forma = next(s for s in slide.shapes if s.name.startswith('Retângulo'))
    assert '<a:alpha val="40000"' in _xml(forma)


def test_formas_nao_herdam_estilo_do_tema_com_sombra_e_contorno(tmp_path):
    """Sem tirar o estilo, o PowerPoint poe sombra e contorno que o PDF nao tem."""
    _, slide = _converte(tmp_path, lambda p: p.draw_rect(fitz.Rect(40, 40, 200, 140), color=None, fill=AZUL))
    forma = next(s for s in slide.shapes if s.name.startswith('Retângulo'))
    assert forma._element.find(qn('p:style')) is None
    assert '<a:effectLst/>' in _xml(forma) and '<a:ln><a:noFill/></a:ln>' in _xml(forma).replace(' xmlns', ' xmlns')


def test_retangulo_do_tamanho_da_pagina_vira_o_fundo_do_slide(tmp_path):
    def desenha(page):
        page.draw_rect(page.rect, color=None, fill=(0.8, 0.8, 0.6))
        page.draw_rect(fitz.Rect(40, 40, 200, 140), color=None, fill=AZUL)
    resultado, slide = _converte(tmp_path, desenha)
    assert str(slide.background.fill.fore_color.rgb) == 'CCCC99'
    assert resultado.formas == 1 and [s.name for s in slide.shapes] == ['Retângulo 1']


# ── imagens ──────────────────────────────────────────────────────────────────

def test_imagem_com_transparencia_vira_figura_solta_com_alfa(tmp_path):
    def desenha(page):
        page.insert_image(fitz.Rect(40, 40, 160, 120), stream=_png((255, 0, 0, 100)))
    resultado, slide = _converte(tmp_path, desenha)
    figuras = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert resultado.imagens == 1 and len(figuras) == 1
    assert figuras[0].left == 40 * 12700 and figuras[0].width == 120 * 12700
    assert fitz.Pixmap(figuras[0].image.blob).alpha                       # a transparencia sobreviveu


def test_ordem_de_empilhamento_do_pdf_e_mantida(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(20, 20, 140, 120), color=None, fill=VERMELHO)                    # 1o: forma
        page.insert_image(fitz.Rect(80, 60, 200, 140), stream=_png((0, 0, 255, 255)))             # 2o: imagem por cima
        page.draw_rect(fitz.Rect(120, 100, 240, 180), color=None, fill=VERDE)                     # 3o: forma por cima
    _, slide = _converte(tmp_path, desenha)
    assert _tipos(slide) == [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.AUTO_SHAPE]


# ── seguranca: o que nao fica igual ao PDF continua imagem ───────────────────

def _pdf_com_recorte(tmp_path, nome, desenha, recorte: bytes | None):
    """PDF de 300x200 pt desenhado por `desenha`, com o conteudo inteiro dentro de um recorte
    (clip) escrito direto no fluxo, ja que o PyMuPDF nao expoe isso. Sem `recorte`, e o controle."""
    origem = tmp_path / f'{nome}.pdf'
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    desenha(page)
    page.clean_contents()
    if recorte is not None:
        xref = page.get_contents()[0]
        doc.update_stream(xref, b'q ' + recorte + b' W n\n' + doc.xref_stream(xref) + b'\nQ')
    doc.save(str(origem))
    doc.close()
    return origem


_CIRCULO = b'200 100 m 200 155 155 200 100 200 c 45 200 0 155 0 100 c 0 45 45 0 100 0 c 155 0 200 45 200 100 c h'
_JANELA = b'100 100 100 100 re'


def test_foto_cortada_em_canto_redondo_nao_vira_figura_quadrada(tmp_path):
    def desenha(page):
        page.insert_image(fitz.Rect(20, 0, 220, 200), stream=_png((255, 140, 0, 255)))
    controle = _pdf_com_recorte(tmp_path, 'sem_clip', desenha, None)
    assert converter_pdf_para_pptx(str(controle), str(tmp_path / 'a.pptx'), MODO_EDITAVEL).imagens == 1
    cortada = _pdf_com_recorte(tmp_path, 'com_clip', desenha, _CIRCULO)
    resultado = converter_pdf_para_pptx(str(cortada), str(tmp_path / 'b.pptx'), MODO_EDITAVEL)
    assert resultado.imagens == 0, 'a imagem recortada nao pode virar um retangulo solto'
    assert resultado.paginas_com_parte_em_imagem == 1


def test_forma_cortada_por_janela_retangular_vira_so_a_parte_visivel(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(10, 10, 290, 190), color=None, fill=VERMELHO)
    controle = _pdf_com_recorte(tmp_path, 'sem_clip', desenha, None)
    assert converter_pdf_para_pptx(str(controle), str(tmp_path / 'a.pptx'), MODO_EDITAVEL).formas == 1
    cortada = _pdf_com_recorte(tmp_path, 'com_clip', desenha, _JANELA)
    resultado = converter_pdf_para_pptx(str(cortada), str(tmp_path / 'b.pptx'), MODO_EDITAVEL)
    forma = next(s for s in Presentation(str(tmp_path / 'b.pptx')).slides[0].shapes)
    assert resultado.formas == 1 and resultado.paginas_com_parte_em_imagem == 0
    # a janela e (100,0)-(200,100) e o retangulo comeca em y=10: visivel = (100,10)-(200,100)
    assert (forma.left, forma.top, forma.width, forma.height) == (100 * 12700, 10 * 12700, 100 * 12700, 90 * 12700)


def test_forma_cortada_por_recorte_redondo_vira_imagem_e_nao_forma_errada(tmp_path):
    def desenha(page):
        page.draw_rect(fitz.Rect(10, 10, 290, 190), color=None, fill=VERMELHO)
    cortada = _pdf_com_recorte(tmp_path, 'redondo', desenha, _CIRCULO)
    resultado = converter_pdf_para_pptx(str(cortada), str(tmp_path / 'r.pptx'), MODO_EDITAVEL)
    slide = Presentation(str(tmp_path / 'r.pptx')).slides[0]
    assert resultado.formas == 0, 'um retangulo inteiro nao pode aparecer onde so cabia um circulo'
    assert resultado.paginas_com_parte_em_imagem == 1 and MSO_SHAPE_TYPE.PICTURE in _tipos(slide)


def test_pagina_com_formas_demais_fica_como_imagem(tmp_path):
    def desenha(page):
        for i in range(850):                                              # acima do limite: e texto vetorizado ou mapa
            page.draw_rect(fitz.Rect(i % 40 * 9, i // 40 * 9, i % 40 * 9 + 6, i // 40 * 9 + 6), color=None, fill=VERMELHO)
    resultado, slide = _converte(tmp_path, desenha, largura=400, altura=250)
    assert resultado.formas == 0 and _tipos(slide) == [MSO_SHAPE_TYPE.PICTURE]


def test_erro_inesperado_na_conferencia_cai_para_imagem_sem_quebrar(tmp_path, monkeypatch):
    def quebra(*args, **kwargs):
        raise RuntimeError('pdf esquisito')
    monkeypatch.setattr(pdf_to_pptx, 'redesenhar', quebra)
    resultado, slide = _converte(tmp_path, lambda p: p.draw_rect(fitz.Rect(40, 40, 200, 140), color=None, fill=AZUL))
    assert resultado.formas == 0 and _tipos(slide) == [MSO_SHAPE_TYPE.PICTURE]


def test_paginas_de_tamanhos_diferentes_com_formas_ficam_dentro_do_slide(tmp_path):
    origem = tmp_path / 'm.pdf'
    doc = fitz.open()
    doc.new_page(width=600, height=400).draw_rect(fitz.Rect(50, 50, 300, 200), color=None, fill=AZUL)
    doc.new_page(width=300, height=500).draw_rect(fitz.Rect(20, 20, 280, 480), color=None, fill=VERMELHO)
    doc.save(str(origem))
    doc.close()
    destino = tmp_path / 'm.pptx'
    converter_pdf_para_pptx(str(origem), str(destino), MODO_EDITAVEL)
    apresentacao = Presentation(str(destino))
    for slide in apresentacao.slides:
        for forma in slide.shapes:
            assert forma.left >= 0 and forma.left + forma.width <= apresentacao.slide_width + 2
            assert forma.top >= 0 and forma.top + forma.height <= apresentacao.slide_height + 2


# ── a conferencia por redesenho ──────────────────────────────────────────────

def _pix(cor, desenho=None):
    doc = fitz.open()
    page = doc.new_page(width=200, height=150)
    page.draw_rect(page.rect, color=None, fill=cor)
    if desenho:
        desenho(page)
    return page.get_pixmap(dpi=60, alpha=False)


def test_conferencia_paginas_iguais_nao_diferem():
    assert pixels_diferentes(_pix((1, 1, 1)), _pix((1, 1, 1))) == 0


def test_conferencia_ve_mancha_larga_mas_ignora_ruido_de_borda():
    cheia = _pix((1, 1, 1), lambda p: p.draw_rect(fitz.Rect(40, 40, 140, 110), color=None, fill=(0, 0, 0)))
    assert pixels_diferentes(_pix((1, 1, 1)), cheia) > 1000
    deslocada = _pix((1, 1, 1), lambda p: p.draw_rect(fitz.Rect(40.3, 40, 140.3, 110), color=None, fill=(0, 0, 0)))
    assert pixels_diferentes(cheia, deslocada) < 30            # meio pixel de diferenca nao reprova
