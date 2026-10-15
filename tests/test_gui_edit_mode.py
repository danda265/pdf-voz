from __future__ import annotations

import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import fitz

from gui import MainWindow


def test_edit_mode_renders_first_page(qtbot, sample_pdf):
    window = MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(sample_pdf)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)
    assert window.edit_session is not None
    assert window.stack.currentIndex() == 1
    assert len(window.edit_scroll.canvas.spans) > 0
    window.close()


def test_documento_do_argv_aceita_pdf_existente(sample_pdf):
    from gui import _documento_do_argv
    assert _documento_do_argv(['app.exe', str(sample_pdf)]) == str(sample_pdf.resolve())


def test_documento_do_argv_ignora_flags_e_arquivo_inexistente(tmp_path):
    from gui import _documento_do_argv
    assert _documento_do_argv(['app.exe']) is None
    assert _documento_do_argv(['app.exe', '--debug']) is None
    assert _documento_do_argv(['app.exe', str(tmp_path / 'nao_existe.pdf')]) is None


def test_documento_do_argv_ignora_extensao_nao_suportada(tmp_path):
    from gui import _documento_do_argv
    outro = tmp_path / 'planilha.xlsx'
    outro.write_bytes(b'x')
    assert _documento_do_argv(['app.exe', str(outro)]) is None


def test_app_abre_com_documento_passado_na_linha_de_comando(qtbot, sample_pdf, monkeypatch):
    """Duplo clique no Explorer entrega o caminho por argv — o app tem de carregar o arquivo."""
    import gui

    carregados = []
    monkeypatch.setattr(gui.MainWindow, '_load_file', lambda self, path: carregados.append(path))
    window = gui.MainWindow()
    qtbot.addWidget(window)
    documento = gui._documento_do_argv(['PDF VOZ.exe', str(sample_pdf)])
    if documento:
        window._load_file(documento)
    window.close()
    assert carregados == [str(sample_pdf.resolve())]


def test_fluxo_completo_editar_e_salvar_pela_interface(qtbot, sample_pdf, monkeypatch):
    """Percorre o caminho real: entra no modo edição, confirma uma edição e salva.

    Existe porque um NameError em _on_popup_confirmed passou por todos os testes de
    unidade e derrubava o app inteiro: no PyQt6 uma exceção não tratada dentro de um
    slot aborta o processo, levando junto as edições não salvas."""
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    salvos = []
    monkeypatch.setattr(gui_mod.QMessageBox, 'information',
                        lambda *a, **k: salvos.append(a[-1]))
    monkeypatch.setattr(gui_mod.QMessageBox, 'critical',
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError(f'erro na UI: {a}')))
    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)

    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(sample_pdf)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)
    assert window.edit_session is not None

    spans = spans_for_page(window.edit_session.doc[0])
    alvo = next(s for s in spans if s.text.strip() == 'Ola mundo')

    # passa pelo caminho real: abrir a caixa no trecho e então confirmar
    window._on_span_clicked(alvo)
    window._on_popup_confirmed([alvo], ['Texto trocado'], FormatOverride(), (0.0, 0.0))
    assert 'Texto trocado' in window.edit_session.doc[0].get_text()
    assert 'Edição aplicada' in window.status_label.text()

    # e salva
    window._save_edited_copy()
    assert salvos, 'não apareceu confirmação de salvamento'
    copia = sample_pdf.with_name(f'{sample_pdf.stem}_editado{sample_pdf.suffix}')
    assert copia.exists()
    assert 'Texto trocado' in fitz.open(str(copia))[0].get_text()

    window.edit_mode_button.setChecked(False)
    window.close()


def test_fluxo_de_varios_trechos_e_multilinha_pela_interface(qtbot, sample_pdf, monkeypatch):
    """Edição de vários trechos e texto com quebra — o caminho que introduziu o bug."""
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    monkeypatch.setattr(gui_mod.QMessageBox, 'critical',
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError(f'erro na UI: {a}')))
    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)
    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(sample_pdf)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    spans = spans_for_page(window.edit_session.doc[0])
    window._on_spans_selected(spans)
    window._on_popup_confirmed(spans, ['linha A\nlinha B'] + ['x'] * (len(spans) - 1),
                               FormatOverride(alinhamento='centro'), (0.0, 0.0))

    texto = window.edit_session.doc[0].get_text()
    assert 'linha A' in texto and 'linha B' in texto
    assert 'Edição aplicada' in window.status_label.text()
    window.edit_mode_button.setChecked(False)
    window.close()


def test_erro_inesperado_nao_derruba_o_app(qtbot, monkeypatch):
    """A rede de segurança tem de engolir a exceção e manter a janela viva — o PyQt6
    aborta o processo quando uma exceção escapa de um slot."""
    import sys as _sys
    import gui as gui_mod

    mostrados = []
    monkeypatch.setattr(gui_mod.QMessageBox, 'critical',
                        lambda *a, **k: mostrados.append(a[2] if len(a) > 2 else ''))
    hook_anterior = _sys.excepthook
    try:
        window = gui_mod.MainWindow()
        qtbot.addWidget(window)
        gui_mod._instalar_rede_de_seguranca(window)
        assert _sys.excepthook is not hook_anterior

        try:
            raise ValueError('falha simulada dentro de um slot')
        except ValueError:
            _sys.excepthook(*_sys.exc_info())   # não pode levantar nem matar o processo

        assert mostrados and 'falha simulada' in mostrados[0]
        assert 'continua aberto' in mostrados[0]
        window.close()
    finally:
        _sys.excepthook = hook_anterior


def _pdf_duas_paginas(tmp_path):
    caminho = tmp_path / 'duas.pdf'
    d = fitz.open()
    a = d.new_page(width=595, height=842)
    a.insert_text((72, 100), 'PAGINA UM ALVO', fontsize=12, fontname='helv')
    b = d.new_page(width=595, height=842)
    b.insert_text((72, 100), 'PAGINA DOIS INTACTA', fontsize=12, fontname='helv')
    d.save(str(caminho)); d.close()
    return caminho


def test_edicao_vai_para_a_pagina_do_trecho_mesmo_apos_rolar(qtbot, tmp_path, monkeypatch):
    """O documento rola contínuo, então não existe "página atual" para errar: cada trecho
    carrega a própria página. Antes, rolar até outra página fazia a edição cair nela,
    apagando conteúdo alheio."""
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    monkeypatch.setattr(gui_mod.QMessageBox, 'critical',
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError(f'erro: {a}')))
    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)

    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(_pdf_duas_paginas(tmp_path))
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    alvo = next(s for s in spans_for_page(window.edit_session.doc[0]) if 'UM ALVO' in s.text)
    assert alvo.pagina == 0
    window._on_span_clicked(alvo)
    pagina_dois_antes = window.edit_session.doc[1].get_text()

    window._go_edit_page(1)   # rola até a página 2 com a caixa aberta
    window._on_popup_confirmed([alvo], ['TROCADO'], FormatOverride(), (0.0, 0.0))

    assert window.edit_session.doc[1].get_text() == pagina_dois_antes, 'mexeu na página errada'
    assert 'TROCADO' in window.edit_session.doc[0].get_text(), 'não editou a página do trecho'
    window.edit_mode_button.setChecked(False)
    window.close()


def test_selecao_que_mistura_paginas_e_recusada(qtbot, tmp_path, monkeypatch):
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    erros = []
    monkeypatch.setattr(gui_mod.QMessageBox, 'critical', lambda *a, **k: erros.append(a))
    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)

    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(_pdf_duas_paginas(tmp_path))
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    de_p1 = spans_for_page(window.edit_session.doc[0])[0]
    de_p2 = spans_for_page(window.edit_session.doc[1])[0]
    antes = [window.edit_session.doc[i].get_text() for i in (0, 1)]

    window._on_popup_confirmed([de_p1, de_p2], ['A', 'B'], FormatOverride(), (0.0, 0.0))

    assert erros, 'aceitou uma seleção que mistura páginas'
    assert [window.edit_session.doc[i].get_text() for i in (0, 1)] == antes
    window.edit_mode_button.setChecked(False)
    window.close()


def test_apagar_linhas_da_selecao_nao_faz_texto_sumir(qtbot, tmp_path, monkeypatch):
    """Verificado em uso: selecionar 3 trechos e deixar 1 linha apagava os três da página
    e escrevia só a linha digitada — os trechos sem par eram redigidos e esquecidos."""
    import gui as gui_mod

    monkeypatch.setattr(gui_mod.QMessageBox, 'critical',
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError(f'erro: {a}')))
    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)

    caminho = tmp_path / 'tres.pdf'
    d = fitz.open()
    pg = d.new_page(width=595, height=842)
    for y, t in ((100, 'ALFA'), (130, 'BETA'), (160, 'GAMA')):
        pg.insert_text((72, y), t, fontsize=12, fontname='helv')
    d.save(str(caminho)); d.close()

    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(caminho)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    from pdf_editor import spans_for_page
    sel = spans_for_page(window.edit_session.doc[0])
    assert len(sel) == 3
    window._on_spans_selected(sel)
    window.edit_popup.text_input.setPlainText('SO UMA LINHA')
    window.edit_popup._confirm()

    final = window.edit_session.doc[0].get_text()
    assert 'SO UMA LINHA' in final
    # os trechos sem linha correspondente viram remoção INTENCIONAL, nunca sumiço silencioso
    assert window.edit_popup._spans == sel
    window.edit_mode_button.setChecked(False)
    window.close()


def test_sair_com_edicoes_nao_salvas_pergunta_antes_de_descartar(qtbot, sample_pdf, monkeypatch):
    """Verificado em uso: sair do modo edição jogava fora tudo sem uma palavra."""
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    perguntas = []

    def responder(*a, **k):
        perguntas.append(a)
        return gui_mod.QMessageBox.StandardButton.Cancel   # usuário desiste de sair

    monkeypatch.setattr(gui_mod.QMessageBox, 'question', responder)
    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(sample_pdf)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    alvo = spans_for_page(window.edit_session.doc[0])[0]
    window._on_span_clicked(alvo)
    window._on_popup_confirmed([alvo], ['EDITADO'], FormatOverride(), (0.0, 0.0))
    assert window._edicoes_pendentes is True

    window.edit_mode_button.setChecked(False)   # tenta sair

    assert perguntas, 'saiu sem perguntar'
    assert window.edit_session is not None, 'descartou mesmo com o usuário cancelando'
    assert 'EDITADO' in window.edit_session.doc[0].get_text()

    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)
    window.edit_mode_button.setChecked(False)
    assert window.edit_session is None
    window.close()


def test_desfazer_fecha_a_caixa_de_edicao(qtbot, sample_pdf, monkeypatch):
    """Desfazer troca o documento inteiro; os trechos guardados na caixa viram
    coordenadas obsoletas e confirmar depois apagaria o que estivesse ali."""
    import gui as gui_mod
    from pdf_editor import FormatOverride, spans_for_page

    monkeypatch.setattr(gui_mod.QMessageBox, 'question',
                        lambda *a, **k: gui_mod.QMessageBox.StandardButton.Discard)
    window = gui_mod.MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(sample_pdf)
    window.edit_mode_button.setEnabled(True)
    window.edit_mode_button.setChecked(True)

    alvo = next(s for s in spans_for_page(window.edit_session.doc[0])
                if s.text.strip() == 'Ola mundo')
    window._on_span_clicked(alvo)
    window._on_popup_confirmed([alvo], ['PRIMEIRA'], FormatOverride(), (0.0, 0.0))

    novo = next(s for s in spans_for_page(window.edit_session.doc[0])
                if 'PRIMEIRA' in s.text)
    window._on_span_clicked(novo)
    assert window.edit_popup.isVisible() or window.edit_popup._spans == [novo]

    window._edit_undo()

    assert not window.edit_popup.isVisible(), 'a caixa continuou aberta com trechos obsoletos'
    assert 'Ola mundo' in window.edit_session.doc[0].get_text()
    window.edit_mode_button.setChecked(False)
    window.close()


# ── Conversao para PowerPoint ────────────────────────────────────────────────

def _janela_com_pdf(qtbot, caminho):
    window = MainWindow()
    qtbot.addWidget(window)
    window.current_file = str(caminho)
    window._document_loaded('texto', 1, 'PDF', ['texto'])
    return window


def _escolhe_arquivo(monkeypatch, caminho):
    """Faz o seletor de arquivo 'Abrir' devolver `caminho` (ou nada, se for None)."""
    from PyQt6.QtWidgets import QFileDialog
    monkeypatch.setattr(QFileDialog, 'getOpenFileName',
                        staticmethod(lambda *a, **k: (str(caminho) if caminho else '', '')))


def test_ferramentas_do_cabecalho_ficam_disponiveis_sem_arquivo_aberto(qtbot):
    window = MainWindow()
    qtbot.addWidget(window)
    for botao in (window.convert_button, window.edit_mode_button, window.export_button,
                  window.merge_button, window.open_button):
        assert botao.isEnabled()
    assert not window.close_button.isEnabled()        # nao ha o que fechar
    window._document_loaded('texto', 1, 'Word', [])
    assert window.convert_button.isEnabled() and window.edit_mode_button.isEnabled()
    window.close_document()
    assert window.convert_button.isEnabled() and window.edit_mode_button.isEnabled()
    window.close()


def test_converter_sem_pdf_aberto_pede_o_arquivo_e_converte(qtbot, sample_pdf, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QFileDialog
    from pptx import Presentation
    from pdf_to_pptx import MODO_EDITAVEL
    window = MainWindow()
    qtbot.addWidget(window)
    _escolhe_arquivo(monkeypatch, sample_pdf)
    destino = tmp_path / 'saida.pptx'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', staticmethod(lambda *a, **k: (str(destino), '')))
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda texto, d: None)
    window.converter_para_pptx(MODO_EDITAVEL)
    qtbot.waitUntil(lambda: window.converter_thread is None, timeout=10000)
    assert len(Presentation(str(destino)).slides) == 1
    assert window.current_file is None                 # converteu sem abrir o arquivo no leitor
    window.close()


def test_converter_cancelando_o_seletor_nao_faz_nada(qtbot, monkeypatch):
    from pdf_to_pptx import MODO_EDITAVEL
    window = MainWindow()
    qtbot.addWidget(window)
    _escolhe_arquivo(monkeypatch, None)
    iniciou = []
    monkeypatch.setattr(window, '_iniciar_conversao', lambda *a: iniciou.append(a))
    window.converter_para_pptx(MODO_EDITAVEL)
    assert iniciou == [] and window.converter_thread is None
    window.close()


def test_editar_pdf_sem_arquivo_pede_o_pdf_e_entra_no_modo_edicao(qtbot, sample_pdf, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    _escolhe_arquivo(monkeypatch, sample_pdf)
    window.edit_mode_button.click()
    qtbot.waitUntil(lambda: window.edit_session is not None, timeout=10000)
    assert window.edit_mode_button.isChecked() and window.stack.currentIndex() == 1
    assert window.current_file == str(sample_pdf)
    window.close()


def test_editar_pdf_cancelando_o_seletor_desmarca_o_botao(qtbot, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    _escolhe_arquivo(monkeypatch, None)
    window.edit_mode_button.click()
    assert not window.edit_mode_button.isChecked() and window.edit_session is None
    assert window.stack.currentIndex() == 0
    window.close()


def test_editar_pdf_com_word_aberto_pede_um_pdf(qtbot, sample_pdf, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    window.current_file = 'C:/docs/relatorio.docx'
    window._document_loaded('texto', 1, 'Word', [])
    _escolhe_arquivo(monkeypatch, sample_pdf)
    window.edit_mode_button.click()
    qtbot.waitUntil(lambda: window.edit_session is not None, timeout=10000)
    assert window.current_file == str(sample_pdf)
    window.close()


def test_exportar_mp3_sem_arquivo_pede_o_arquivo_e_retoma_depois_de_carregar(qtbot, sample_pdf, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    _escolhe_arquivo(monkeypatch, sample_pdf)
    window.export_mp3()                                   # sem texto: pede o arquivo
    retomou = []
    monkeypatch.setattr(window, 'export_mp3', lambda: retomou.append(True))   # o real falaria com a internet
    qtbot.waitUntil(lambda: bool(retomou), timeout=10000)
    assert window.current_text.strip() and window._depois_de_carregar is None
    window.close()


def test_arquivo_que_nao_abre_descarta_a_tarefa_pendente(qtbot, tmp_path, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    ruim = tmp_path / 'ruim.pdf'
    ruim.write_bytes(b'isto nao e um pdf')
    _escolhe_arquivo(monkeypatch, ruim)
    erros, retomou = [], []
    monkeypatch.setattr(window, '_show_error', erros.append)
    window._pedir_arquivo_e_depois(lambda: retomou.append(True))
    qtbot.waitUntil(lambda: bool(erros), timeout=10000)
    assert window._depois_de_carregar is None and retomou == []
    window.close()


def test_conversao_pelo_app_grava_o_pptx(qtbot, sample_pdf, tmp_path, monkeypatch):
    from pptx import Presentation
    from pdf_to_pptx import MODO_EDITAVEL
    window = _janela_com_pdf(qtbot, sample_pdf)
    avisos = []
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda texto, destino: avisos.append((texto, destino)))
    destino = tmp_path / 'saida.pptx'
    window._iniciar_conversao(MODO_EDITAVEL, str(destino))
    qtbot.waitUntil(lambda: window.converter_thread is None, timeout=10000)
    assert destino.exists() and len(Presentation(str(destino)).slides) == 1
    assert len(avisos) == 1 and 'caixas de texto' in avisos[0][0] and avisos[0][1] == str(destino)
    assert window.convert_button.isEnabled()
    assert window.progress_bar.value() == 100
    window.close()


def test_conversao_com_erro_mostra_a_mensagem_e_libera_o_botao(qtbot, sample_pdf, tmp_path, monkeypatch):
    from pdf_to_pptx import MODO_EDITAVEL
    window = _janela_com_pdf(qtbot, sample_pdf)
    erros = []
    monkeypatch.setattr(window, '_show_error', erros.append)
    window._iniciar_conversao(MODO_EDITAVEL, str(sample_pdf))   # saida = o proprio PDF: recusado
    qtbot.waitUntil(lambda: window.converter_thread is None, timeout=10000)
    assert len(erros) == 1 and 'próprio PDF' in erros[0]
    assert window.convert_button.isEnabled()
    window.close()


def test_conversao_cancelada_nao_grava_nada(qtbot, multipage_pdf, tmp_path):
    from gui import ConverterThread
    from pdf_to_pptx import MODO_IMAGEM
    destino = tmp_path / 'saida.pptx'
    thread = ConverterThread(str(multipage_pdf), str(destino), MODO_IMAGEM)
    concluidos, erros = [], []
    thread.done.connect(concluidos.append)
    thread.error.connect(erros.append)
    thread.cancelar()           # antes de comecar: o primeiro aviso de progresso ja interrompe
    thread.start()
    assert thread.wait(10000)
    assert concluidos == [] and erros == []
    assert not destino.exists() and list(tmp_path.glob('*.parcial')) == []


def test_fechar_a_janela_espera_a_conversao_terminar(qtbot, multipage_pdf, tmp_path, monkeypatch):
    from pdf_to_pptx import MODO_IMAGEM
    window = _janela_com_pdf(qtbot, multipage_pdf)
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda *a: None)
    window._iniciar_conversao(MODO_IMAGEM, str(tmp_path / 'saida.pptx'))
    window.close()
    assert window.converter_thread is None or not window.converter_thread.isRunning()
    assert list((tmp_path).glob('*.parcial')) == []


def test_edicoes_nao_salvas_pedem_confirmacao(qtbot, sample_pdf, monkeypatch):
    from PyQt6.QtWidgets import QMessageBox
    window = _janela_com_pdf(qtbot, sample_pdf)
    window._edicoes_pendentes = True
    perguntas = []
    monkeypatch.setattr(QMessageBox, 'question',
                        staticmethod(lambda *a, **k: perguntas.append(a) or QMessageBox.StandardButton.No))
    iniciou = []
    monkeypatch.setattr(window, '_iniciar_conversao', lambda *a: iniciou.append(a))
    window.converter_para_pptx('editavel')
    assert len(perguntas) == 1 and iniciou == []
    window.close()


def test_mensagem_da_conversao_conta_formas_imagens_e_o_que_ficou_como_imagem(qtbot, tmp_path, monkeypatch):
    import fitz
    from pdf_to_pptx import MODO_EDITAVEL
    origem = tmp_path / 'forma.pdf'
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    page.draw_rect(fitz.Rect(20, 20, 200, 120), color=None, fill=(1, 0, 0))
    page.insert_text((40, 70), 'texto', fontsize=14, fontname='helv', color=(1, 1, 1))
    doc.save(str(origem))
    doc.close()
    window = _janela_com_pdf(qtbot, origem)
    avisos = []
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda texto, destino: avisos.append(texto))
    window._iniciar_conversao(MODO_EDITAVEL, str(tmp_path / 'saida.pptx'))
    qtbot.waitUntil(lambda: window.converter_thread is None, timeout=10000)
    assert '1 forma e 0 imagens separadas' in avisos[0] and '1 caixas de texto' in avisos[0]
    assert 'ficaram como imagem' not in avisos[0]              # nada ficou de fora nesta pagina
    window.close()


def test_menu_converter_tem_os_tres_modos_e_o_ocr_dispara_o_modo_ocr(qtbot, monkeypatch):
    from pdf_to_pptx import MODO_EDITAVEL, MODO_IMAGEM, MODO_OCR
    window = MainWindow()
    qtbot.addWidget(window)
    acoes = window.convert_button.menu().actions()
    assert len(acoes) == 3 and any('OCR' in a.text() for a in acoes)
    pedidos = []
    monkeypatch.setattr(window, 'converter_para_pptx', pedidos.append)
    window.act_pptx_editavel.trigger()
    window.act_pptx_ocr.trigger()
    window.act_pptx_imagem.trigger()
    assert pedidos == [MODO_EDITAVEL, MODO_OCR, MODO_IMAGEM]
    window.close()


def test_mensagem_do_modo_ocr_conta_os_paragrafos_e_avisa_que_o_ocr_pode_errar(qtbot, monkeypatch):
    from pdf_to_pptx import ResultadoConversao
    window = MainWindow()
    qtbot.addWidget(window)
    avisos = []
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda texto, destino: avisos.append(texto))
    window._conversao_concluida(ResultadoConversao(
        'x.pptx', 2, paginas_com_parte_em_imagem=2, paginas_com_ocr=2, paragrafos_ocr=7, linhas_ocr_ignoradas=3))
    texto = avisos[0]
    assert '7 parágrafos reconhecidos por OCR em 2 páginas' in texto and 'revise o texto' in texto
    assert '3 linhas ficaram na imagem' in texto
    assert 'não dá para separar' not in texto           # a imagem das paginas com OCR nao conta como "sobrou"
    window.close()


def test_mensagem_com_uma_pagina_so_imagem_e_outra_com_ocr_separa_os_avisos(qtbot, monkeypatch):
    from pdf_to_pptx import ResultadoConversao
    window = MainWindow()
    qtbot.addWidget(window)
    avisos = []
    monkeypatch.setattr(window, '_avisar_conversao_pronta', lambda texto, destino: avisos.append(texto))
    window._conversao_concluida(ResultadoConversao(
        'x.pptx', 3, paginas_com_parte_em_imagem=3, paginas_com_ocr=1, paragrafos_ocr=1, linhas_ocr_ignoradas=1))
    assert '1 parágrafo reconhecido por OCR em 1 página' in avisos[0]
    assert 'Em 2 página(s), partes que não dá para separar' in avisos[0]
    window.close()
