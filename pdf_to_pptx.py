"""Conversao de PDF para PowerPoint (.pptx), sem dependencia de interface.

Dois modos:
  - MODO_IMAGEM: cada pagina vira um slide com a pagina inteira como imagem. Fiel ao PDF, mas
    nada nele e editavel.
  - MODO_EDITAVEL: o slide tem de fundo a pagina SEM o texto horizontal, e o texto volta como
    caixas de texto do PowerPoint, na mesma posicao, fonte, tamanho e cor. Graficos, fotos e
    texto girado ficam no fundo.

O PDF original nunca e alterado: as remocoes de texto acontecem so na copia em memoria.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from statistics import median
from typing import Callable, List, Optional, Tuple

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from pdf_editor import familia_legivel

MODO_IMAGEM = 'imagem'
MODO_EDITAVEL = 'editavel'

# Resolucao do fundo de cada slide. O modo editavel usa menos porque o texto sai da imagem.
_DPI = {MODO_IMAGEM: 200, MODO_EDITAVEL: 150}
_LADO_MAXIMO_PX = 4000          # trava de memoria para paginas enormes (cartazes)
_EMU_POR_PT = 12700
_SLIDE_MAX_PT = 4000            # o PowerPoint nao aceita slide acima de 56 pol (4032 pt)
_SLIDE_MIN_PT = 72              # ... nem abaixo de 1 pol
_CONTROLE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

ProgressCallback = Optional[Callable[[int, int], None]]


class ConversaoError(Exception):
    """Falha esperada da conversao, com mensagem que pode ir direto para o usuario."""


class ConversaoCancelada(Exception):
    """O chamador interrompeu a conversao (levantada de dentro do progress_callback).

    Nenhum arquivo de saida e criado: o pptx so e gravado depois da ultima pagina."""


@dataclass
class ResultadoConversao:
    destino: str
    paginas: int
    caixas_de_texto: int = 0
    paginas_sem_texto_editavel: int = 0   # no modo editavel: paginas que ficaram so como imagem


@dataclass
class _Linha:
    bbox: Tuple[float, float, float, float]
    baseline: float
    spans: List[dict]


def _limpa(texto: str) -> str:
    # caracteres de controle quebram o XML do pptx
    return _CONTROLE.sub('', texto)


def _tamanho_do_slide(pagina: fitz.Rect) -> Tuple[float, float]:
    """Tamanho do slide em pontos: o da primeira pagina, ajustado ao que o PowerPoint aceita."""
    largura, altura = pagina.width, pagina.height
    escala = 1.0
    if max(largura, altura) > _SLIDE_MAX_PT:
        escala = _SLIDE_MAX_PT / max(largura, altura)
    if min(largura, altura) * escala < _SLIDE_MIN_PT:
        escala = _SLIDE_MIN_PT / min(largura, altura)
    return largura * escala, altura * escala


def _encaixe(pagina: fitz.Rect, slide: Tuple[float, float]) -> Tuple[float, float, float]:
    """Escala e deslocamento (em pt do slide) que colocam a pagina inteira dentro do slide."""
    escala = min(slide[0] / pagina.width, slide[1] / pagina.height)
    return (escala,
            (slide[0] - pagina.width * escala) / 2,
            (slide[1] - pagina.height * escala) / 2)


def _emu(valor_pt: float) -> Emu:
    return Emu(int(round(valor_pt * _EMU_POR_PT)))


def _imagem_da_pagina(page: fitz.Page, dpi: int) -> BytesIO:
    """Renderiza a pagina. JPEG quando ha foto (PNG ficaria enorme), PNG nos demais."""
    maior_lado_pt = max(page.rect.width, page.rect.height)
    dpi = max(36, min(dpi, int(_LADO_MAXIMO_PX / maior_lado_pt * 72)))
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    if page.get_images():
        return BytesIO(pix.tobytes('jpeg', jpg_quality=88))
    return BytesIO(pix.tobytes('png'))


def _caixas_invisiveis(page: fitz.Page) -> List[fitz.Rect]:
    """Areas com texto invisivel — a camada de OCR dos PDFs escaneados.

    Esse texto fica por cima da imagem do scan sem aparecer; virar caixa de texto o tornaria
    visivel e duplicado."""
    try:
        trace = page.get_texttrace()
    except Exception:
        return []
    return [fitz.Rect(t['bbox']) for t in trace if t.get('type') == 3]


def _linhas_editaveis(page: fitz.Page) -> List[List[_Linha]]:
    """Blocos de texto horizontal e visivel da pagina, cada um com suas linhas."""
    invisiveis = _caixas_invisiveis(page)
    blocos: List[List[_Linha]] = []
    for bloco in page.get_text('dict').get('blocks', []):
        if bloco.get('type') != 0:
            continue
        linhas: List[_Linha] = []
        for linha in bloco.get('lines', []):
            direcao = linha.get('dir', (1, 0))
            if abs(direcao[1]) > 0.01 or direcao[0] < 0:   # girado ou de tras para frente
                continue
            spans = [s for s in linha.get('spans', []) if _limpa(s.get('text', '')).strip()]
            if not spans:
                continue
            bbox = fitz.Rect(linha['bbox'])
            centro = fitz.Point((bbox.x0 + bbox.x1) / 2, (bbox.y0 + bbox.y1) / 2)
            if any(centro in area for area in invisiveis):
                continue
            linhas.append(_Linha(tuple(bbox), spans[0]['origin'][1], spans))
        if linhas:
            blocos.append(linhas)
    return blocos


def _tamanho_da_linha(linha: _Linha) -> float:
    return max(s['size'] for s in linha.spans)


def _grupos_de_linhas(linhas: List[_Linha]) -> List[List[_Linha]]:
    """Linhas que alinham na esquerda e tem o mesmo tamanho formam um paragrafo so (uma caixa).

    As demais — centralizadas, ou de tamanhos diferentes como um titulo colado no subtitulo —
    ficam cada uma na sua caixa: o espaçamento exato de uma caixa so serve a linhas parecidas,
    e em linhas desiguais empurraria a primeira para baixo."""
    if len(linhas) > 1:
        alinhadas = max(l.bbox[0] for l in linhas) - min(l.bbox[0] for l in linhas) <= 2.0
        tamanhos = [_tamanho_da_linha(l) for l in linhas]
        if alinhadas and max(tamanhos) <= min(tamanhos) * 1.15:
            return [linhas]
    return [[l] for l in linhas]


def _aplica_span(run, span: dict, escala: float) -> None:
    run.text = _limpa(span['text'])
    fonte = run.font
    fonte.name = familia_legivel(span.get('font', '')) or 'Arial'
    fonte.size = Pt(max(1.0, span['size'] * escala))
    flags = span.get('flags', 0)
    nome = (span.get('font') or '').lower()
    fonte.bold = bool(flags & 16) or 'bold' in nome
    fonte.italic = bool(flags & 2) or 'italic' in nome or 'oblique' in nome
    cor = span.get('color', 0)
    fonte.color.rgb = RGBColor((cor >> 16) & 255, (cor >> 8) & 255, cor & 255)


def _adiciona_caixa(slide, grupo: List[_Linha], escala: float, dx: float, dy: float) -> None:
    x0 = min(l.bbox[0] for l in grupo)
    y0 = min(l.bbox[1] for l in grupo)
    x1 = max(l.bbox[2] for l in grupo)
    y1 = max(l.bbox[3] for l in grupo)
    caixa = slide.shapes.add_textbox(_emu(dx + x0 * escala), _emu(dy + y0 * escala),
                                     _emu((x1 - x0) * escala), _emu((y1 - y0) * escala))
    quadro = caixa.text_frame
    quadro.margin_left = quadro.margin_right = quadro.margin_top = quadro.margin_bottom = 0
    quadro.word_wrap = False           # as quebras de linha sao as do PDF, nao do PowerPoint
    quadro.auto_size = MSO_AUTO_SIZE.NONE
    quadro.vertical_anchor = MSO_ANCHOR.TOP
    passos = [b.baseline - a.baseline for a, b in zip(grupo, grupo[1:])]
    for indice, linha in enumerate(grupo):
        paragrafo = quadro.paragraphs[0] if indice == 0 else quadro.add_paragraph()
        if passos:
            paragrafo.line_spacing = Pt(median(passos) * escala)
        for span in linha.spans:
            _aplica_span(paragrafo.add_run(), span, escala)


def _remove_o_texto_horizontal(page: fitz.Page, blocos: List[List[_Linha]]) -> None:
    """Apaga da pagina (em memoria) o texto que vai virar caixa, sem encostar em fotos nem desenhos."""
    for bloco in blocos:
        for linha in bloco:
            rect = fitz.Rect(linha.bbox)
            folga = rect.height * 0.12     # nao invade a linha vizinha
            rect = fitz.Rect(rect.x0, rect.y0 + folga, rect.x1, rect.y1 - folga)
            page.add_redact_annot(rect, fill=False)
    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                          graphics=fitz.PDF_REDACT_LINE_ART_NONE)


def converter_pdf_para_pptx(origem: str, destino: str, modo: str = MODO_EDITAVEL,
                            progress_callback: ProgressCallback = None) -> ResultadoConversao:
    if modo not in (MODO_IMAGEM, MODO_EDITAVEL):
        raise ConversaoError('Modo de conversão desconhecido.')
    origem_p, destino_p = Path(origem), Path(destino)
    if origem_p.suffix.lower() != '.pdf' or not origem_p.is_file():
        raise ConversaoError('Abra um PDF válido para converter.')
    if destino_p.resolve() == origem_p.resolve():
        raise ConversaoError('O arquivo de saída não pode ser o próprio PDF.')

    try:
        doc = fitz.open(str(origem_p))
    except Exception as exc:
        raise ConversaoError(f'Não foi possível abrir o PDF: {exc}') from exc

    with doc:
        if doc.needs_pass:
            raise ConversaoError('O PDF está protegido por senha.')
        total = doc.page_count
        if total == 0:
            raise ConversaoError('O PDF não tem páginas.')

        slide_pt = _tamanho_do_slide(doc[0].rect)
        apresentacao = Presentation()
        apresentacao.slide_width = _emu(slide_pt[0])
        apresentacao.slide_height = _emu(slide_pt[1])
        layout_vazio = apresentacao.slide_layouts[6]
        resultado = ResultadoConversao(str(destino_p), total)

        for indice in range(total):
            page = doc[indice]
            escala, dx, dy = _encaixe(page.rect, slide_pt)
            blocos: List[List[_Linha]] = []
            # pagina girada: as coordenadas do texto nao batem com a imagem, fica so a imagem
            if modo == MODO_EDITAVEL and page.rotation == 0:
                blocos = _linhas_editaveis(page)
            if blocos:
                _remove_o_texto_horizontal(page, blocos)
            elif modo == MODO_EDITAVEL:
                resultado.paginas_sem_texto_editavel += 1

            dados = _imagem_da_pagina(page, _DPI[modo])
            slide = apresentacao.slides.add_slide(layout_vazio)
            slide.shapes.add_picture(dados, _emu(dx), _emu(dy),
                                     _emu(page.rect.width * escala),
                                     _emu(page.rect.height * escala))
            for bloco in blocos:
                for grupo in _grupos_de_linhas(bloco):
                    _adiciona_caixa(slide, grupo, escala, dx, dy)
                    resultado.caixas_de_texto += 1
            if progress_callback:
                progress_callback(indice + 1, total)

    temporario = destino_p.with_name(destino_p.name + '.parcial')
    try:
        apresentacao.save(str(temporario))
        os.replace(temporario, destino_p)
    except OSError as exc:
        raise ConversaoError(f'Não foi possível salvar o arquivo (ele está aberto em outro programa?): {exc}') from exc
    finally:
        if temporario.exists():
            temporario.unlink()
    return resultado
