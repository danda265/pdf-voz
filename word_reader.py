from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from utils import clean_text


@dataclass
class WordText:
    text: str
    total_pages: int = 1


class WordReader:
    def read(self, path: str) -> WordText:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix == '.docx':
            return self._read_docx(file_path)
        if suffix == '.doc':
            return self._read_doc(file_path)
        raise ValueError('Formato Word nao suportado.')

    def _read_docx(self, path: Path) -> WordText:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError('Para ler Word, instale: pip install python-docx') from exc

        document = Document(str(path))
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(' | '.join(cells))
        return WordText(clean_text('\n\n'.join(parts)), 1)

    def _read_doc(self, path: Path) -> WordText:
        """Converte .doc para .docx usando Microsoft Word, quando disponivel no Windows."""
        try:
            import win32com.client
        except ImportError as exc:
            raise RuntimeError(
                'Arquivos .doc antigos exigem Microsoft Word e pywin32. Instale: pip install pywin32, ou salve como .docx.'
            ) from exc

        with TemporaryDirectory() as tmp:
            output = Path(tmp) / f'{path.stem}.docx'
            word = win32com.client.Dispatch('Word.Application')
            word.Visible = False
            try:
                doc = word.Documents.Open(str(path.resolve()))
                doc.SaveAs(str(output), FileFormat=16)
                doc.Close(False)
            finally:
                word.Quit()
            return self._read_docx(output)
