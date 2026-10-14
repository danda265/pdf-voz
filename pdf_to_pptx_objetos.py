"""Formas e imagens de uma pagina de PDF, recriadas como objetos do PowerPoint.

O PDF guarda o desenho como retangulos, curvas e imagens posicionadas. Aqui cada um vira um
objeto proprio do slide (que da para mover, recolorir, trocar ou apagar), na mesma ordem de
empilhamento do PDF. O que o PDF nao descreve como forma — degrades, sombreados — e devolvido
como "recorte" para o chamador cortar da imagem da pagina.

Regra de seguranca: antes de usar, a conversao e conferida. A pagina e redesenhada so com os
objetos que seriam criados e comparada, pixel a pixel, com o original (`pixels_diferentes`).
Foto cortada em canto redondo, transparencia em grupo, forma mal lida: qualquer coisa que mude
a aparencia reprova a pagina, que continua como imagem de fundo. Melhor um fundo achatado do
que um slide diferente do PDF.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Dict, List, Optional, Sequence, Tuple

import fitz
from PIL import Image, ImageChops, ImageFilter
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

_EMU_POR_PT = 12700
_TOL = 0.6              # tolerancia geometrica, em pontos
_MAX_FORMAS = 800       # acima disso e provavel que seja texto vetorizado ou mapa: fica como imagem
_CAPS = {0: 'flat', 1: 'rnd', 2: 'sq'}


# ── Analise da pagina ────────────────────────────────────────────────────────

@dataclass
class Forma:
    ordem: int
    desenho: dict


@dataclass
class Figura:
    ordem: int
    caixa: fitz.Rect
    dados: bytes
    extensao: str


@dataclass
class Recorte:
    """Parte da pagina que o PDF nao descreve como forma (degrade/sombreado): sai da imagem."""
    ordem: int
    caixa: fitz.Rect


@dataclass
class AnaliseDaPagina:
    formas: List[Forma] = field(default_factory=list)
    figuras: List[Figura] = field(default_factory=list)
    recortes: List[Recorte] = field(default_factory=list)
    formas_ok: bool = False     # todas as formas da pagina sao de tipos que sabemos recriar
    figuras_ok: bool = False    # todas as imagens da pagina deram para extrair


def _perto(a, b, tol: float = _TOL) -> bool:
    return abs(a[0] - b[0]) <= tol and abs(a[1] - b[1]) <= tol


def _bytes_da_imagem(doc: fitz.Document, xref: int) -> Optional[Tuple[bytes, str]]:
    """Bytes de uma imagem do PDF prontos para o PowerPoint, ou None se nao der para extrair."""
    try:
        if doc.xref_get_key(xref, 'ImageMask')[1] == 'true':
            return None   # mascara de recorte: a cor vem do contexto, nao da imagem
        info = doc.extract_image(xref)
        if not info:
            return None
        sem_alpha = not info.get('smask')
        if sem_alpha and info.get('ext') in ('png', 'jpeg', 'jpg') and info.get('colorspace') in (1, 3):
            return info['image'], 'png' if info['ext'] == 'png' else 'jpg'
        pix = fitz.Pixmap(doc, xref)
        if pix.colorspace is not None and pix.colorspace.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        if pix.alpha:
            pix = fitz.Pixmap(pix, 0)           # tira o alfa proprio; vale a mascara
        if info.get('smask'):
            mascara = fitz.Pixmap(doc, info['smask'])
            if (mascara.width, mascara.height) != (pix.width, pix.height):
                return None
            pix = fitz.Pixmap(pix, mascara)
        return pix.tobytes('png'), 'png'
    except Exception:
        return None


def _clip_retangular(clip: dict) -> bool:
    itens = clip.get('items') or []
    if len(itens) == 1 and itens[0][0] in ('re', 'qu'):
        return True
    return len(itens) in (4, 5) and all(i[0] == 'l' for i in itens) and all(
        abs(i[1][0] - i[2][0]) <= _TOL or abs(i[1][1] - i[2][1]) <= _TOL for i in itens)


def _dentro(caixa: fitz.Rect, janela: fitz.Rect, margem: float) -> bool:
    return (caixa.x0 >= janela.x0 + margem - _TOL and caixa.y0 >= janela.y0 + margem - _TOL
            and caixa.x1 <= janela.x1 - margem + _TOL and caixa.y1 <= janela.y1 - margem + _TOL)


def _classifica(d: dict, janelas: Sequence[Tuple[fitz.Rect, bool]]) -> Tuple[Optional[str], Optional[object]]:
    """O que fazer com um desenho que esta dentro de recortes (clip) do PDF.

    Devolve ('forma', desenho) para criar a forma, ('ilha', area) para recortar a area visivel
    da pagina como imagem, ou (None, None) se o recorte o esconde por completo.

    Recorte retangular que nao corta o desenho nao muda nada; um retangulo cheio num recorte
    retangular encolhe para a area visivel. Recorte de outro formato (borda arredondada de
    cartao, por exemplo) que corta o desenho nao tem como ser reproduzido como forma: vira ilha.
    Um recorte de outro formato so e inofensivo se o desenho ficar folgado dentro dele."""
    caixa = fitz.Rect(d['rect'])
    if 's' in d['type']:
        meia = (d.get('width') or 0) / 2
        caixa = fitz.Rect(caixa.x0 - meia, caixa.y0 - meia, caixa.x1 + meia, caixa.y1 + meia)
    if not janelas:
        return 'forma', d
    visivel = fitz.Rect(caixa)
    for janela, _ in janelas:
        visivel = fitz.Rect(max(visivel.x0, janela.x0), max(visivel.y0, janela.y0),
                            min(visivel.x1, janela.x1), min(visivel.y1, janela.y1))
    if visivel.x1 - visivel.x0 <= 0.05 or visivel.y1 - visivel.y0 <= 0.05:
        return None, None
    if all(_dentro(caixa, janela, 0.0 if reta else 1.0) for janela, reta in janelas):
        return 'forma', d
    if all(reta for _, reta in janelas) and d['type'] == 'f' and len(d['items']) == 1             and d['items'][0][0] == 're':
        recortado = dict(d)
        recortado['items'] = [('re', visivel, d['items'][0][2])]
        recortado['rect'] = visivel
        return 'forma', recortado
    return 'ilha', visivel


def analisar_pagina(doc: fitz.Document, page: fitz.Page) -> AnaliseDaPagina:
    """Levanta as formas, as imagens e os recortes da pagina, cada um com sua ordem de desenho.

    Aqui so se checa o que e preciso para criar o objeto; se a conversao fica igual ao PDF quem
    diz e a conferencia por redesenho (`pixels_diferentes`)."""
    analise = AnaliseDaPagina()
    try:
        desenhos = page.get_drawings(extended=True)
        info = page.get_image_info(xrefs=True)
        registro = page.get_bboxlog()
    except Exception:
        return analise

    formas_ok = True
    janelas: List[Tuple[fitz.Rect, bool]] = []    # recortes (clip) abertos, do externo para o interno
    for d in desenhos:
        tipo = d.get('type')
        nivel = d.get('level', 0) or 0
        if tipo == 'clip':
            del janelas[nivel:]
            janelas.append((fitz.Rect(d['scissor']), _clip_retangular(d)))
            continue
        if tipo not in ('f', 's', 'fs') or (d.get('fill') is None and d.get('color') is None):
            continue                              # nao pinta nada
        if not d.get('items') or any(i[0] not in ('l', 'c', 're', 'qu') for i in d['items']):
            formas_ok = False
            continue
        destino, resultado = _classifica(d, janelas[:nivel])
        if destino == 'forma':
            analise.formas.append(Forma(d['seqno'], resultado))
        elif destino == 'ilha':
            analise.recortes.append(Recorte(d['seqno'], resultado))
    analise.formas_ok = formas_ok and len(analise.formas) + len(analise.recortes) <= _MAX_FORMAS

    figuras_ok = True
    posicoes = [i for i, (t, _) in enumerate(registro) if t in ('fill-image', 'fill-shade')]
    if len(posicoes) != len(info):
        figuras_ok = False                        # nao da para casar cada imagem com sua ordem
    else:
        for ordem, im in zip(posicoes, info):
            xref = im.get('xref', 0)
            if xref == 0:                         # degrade/sombreado ou imagem embutida no fluxo
                analise.recortes.append(Recorte(ordem, fitz.Rect(im['bbox'])))
                continue
            t = im['transform']
            achada = _bytes_da_imagem(doc, xref)
            reta = abs(t[1]) < 1e-3 and abs(t[2]) < 1e-3 and t[0] > 0 and t[3] > 0
            if not (reta and achada):
                figuras_ok = False
                continue
            analise.figuras.append(
                Figura(ordem, fitz.Rect(t[4], t[5], t[4] + t[0], t[5] + t[3]), achada[0], achada[1]))
    analise.figuras_ok = figuras_ok
    return analise


# ── Conferencia por redesenho ────────────────────────────────────────────────

def _desenha_forma_no_pdf(destino: fitz.Page, d: dict) -> None:
    forma = destino.new_shape()
    for item in d['items']:
        if item[0] == 'l':
            forma.draw_line(item[1], item[2])
        elif item[0] == 'c':
            forma.draw_bezier(item[1], item[2], item[3], item[4])
        elif item[0] == 're':
            forma.draw_rect(item[1])
        else:
            forma.draw_quad(item[1])
    cap = d.get('lineCap') or 0
    forma.finish(fill=d.get('fill') if 'f' in d['type'] else None,
                 color=d.get('color') if 's' in d['type'] else None,
                 width=d.get('width') or 1, closePath=bool(d.get('closePath')),
                 even_odd=bool(d.get('even_odd')), dashes=d.get('dashes'),
                 lineCap=cap[0] if isinstance(cap, (tuple, list)) else cap,
                 lineJoin=d.get('lineJoin') or 0,
                 fill_opacity=d.get('fill_opacity') if d.get('fill_opacity') is not None else 1,
                 stroke_opacity=d.get('stroke_opacity') if d.get('stroke_opacity') is not None else 1)
    forma.commit()


def redesenhar(tamanho: fitz.Rect, restante: Optional[fitz.Pixmap], itens: Sequence[Tuple[str, object]],
               dpi: int) -> fitz.Pixmap:
    """Renderiza uma pagina feita so do `restante` (fundo) e dos objetos, na ordem recebida.

    `itens` sao ('forma', Forma), ('imagem', Figura) ou ('recorte', (Recorte, Pixmap))."""
    modelo = fitz.open()
    try:
        pagina = modelo.new_page(width=tamanho.width, height=tamanho.height)
        if restante is not None:
            pagina.insert_image(pagina.rect, pixmap=restante)
        for tipo, objeto in itens:
            if tipo == 'forma':
                _desenha_forma_no_pdf(pagina, objeto.desenho)
            elif tipo == 'imagem':
                pagina.insert_image(objeto.caixa, stream=objeto.dados)
            else:
                recorte, pixmap = objeto
                pagina.insert_image(recorte.caixa, pixmap=pixmap)
        return pagina.get_pixmap(dpi=dpi, alpha=False)
    finally:
        modelo.close()


def pixels_diferentes(a: fitz.Pixmap, b: fitz.Pixmap, limiar: int = 24) -> int:
    """Quantos pixels diferem de verdade (acima de `limiar` em algum canal) entre duas paginas.

    Dois cuidados para nao reprovar por ruido: as imagens sao levemente suavizadas (reduzir uma
    imagem grande de dois jeitos da bordas um pouco diferentes) e cada pixel vale a MENOR
    diferenca entre ele e os 8 vizinhos do outro lado (borda recortada ou desenhada com meio
    pixel de desvio). Uma area errada ou uma forma que falta continuam aparecendo."""
    if (a.width, a.height) != (b.width, b.height):
        return a.width * a.height
    suave = ImageFilter.BoxBlur(1)
    imagem_a = Image.frombytes('RGB', (a.width, a.height), a.samples).filter(suave)
    imagem_b = Image.frombytes('RGB', (b.width, b.height), b.samples).filter(suave)
    menor = ImageChops.difference(imagem_a, imagem_b)
    for dx in (-1, 0, 1):
        for dy in (-1, 0, 1):
            if dx or dy:
                menor = ImageChops.darker(menor, ImageChops.difference(imagem_a, ImageChops.offset(imagem_b, dx, dy)))
    miolo = menor.crop((1, 1, a.width - 1, a.height - 1)).convert('L')      # o deslocamento "dobra" a borda
    return sum(miolo.histogram()[limiar + 1:])


# ── Geometria ────────────────────────────────────────────────────────────────

Segmento = Tuple  # ('M', p) | ('L', p) | ('C', p1, p2, p3)


def _caminhos(d: dict) -> List[Tuple[List[Segmento], bool]]:
    """Subcaminhos do desenho como (segmentos, fechado)."""
    caminhos: List[Tuple[List[Segmento], bool]] = []
    atual: Optional[List[Segmento]] = None
    ponto = None
    for item in d['items']:
        tipo = item[0]
        if tipo == 're':
            r = item[1]
            caminhos.append(([('M', (r.x0, r.y0)), ('L', (r.x1, r.y0)), ('L', (r.x1, r.y1)),
                              ('L', (r.x0, r.y1))], True))
            atual = None
            continue
        if tipo == 'qu':
            q = item[1]
            caminhos.append(([('M', tuple(q.ul)), ('L', tuple(q.ur)), ('L', tuple(q.lr)),
                              ('L', tuple(q.ll))], True))
            atual = None
            continue
        inicio = item[1]
        if atual is None or not _perto(ponto, inicio):
            atual = [('M', tuple(inicio))]
            caminhos.append((atual, False))
        if tipo == 'l':
            atual.append(('L', tuple(item[2])))
            ponto = item[2]
        else:
            atual.append(('C', tuple(item[2]), tuple(item[3]), tuple(item[4])))
            ponto = item[4]
    resultado = []
    for segmentos, fechado in caminhos:
        if not fechado:
            ultimo = segmentos[-1][-1]
            fechado = bool(d.get('closePath')) or (len(segmentos) > 2 and _perto(ultimo, segmentos[0][1]))
        resultado.append((segmentos, fechado))
    return resultado


def _pontos(caminhos) -> List[Tuple[float, float]]:
    return [p for segmentos, _ in caminhos for seg in segmentos for p in seg[1:]]


def _ponto_da_curva(p0, c1, c2, p3, t: float) -> Tuple[float, float]:
    u = 1 - t
    return (u ** 3 * p0[0] + 3 * u * u * t * c1[0] + 3 * u * t * t * c2[0] + t ** 3 * p3[0],
            u ** 3 * p0[1] + 3 * u * u * t * c1[1] + 3 * u * t * t * c2[1] + t ** 3 * p3[1])


def _e_retangulo(itens, caixa: fitz.Rect) -> bool:
    if len(itens) == 1 and itens[0][0] == 're':
        return True
    if len(itens) == 1 and itens[0][0] == 'qu':
        q = itens[0][1]
        return abs(q.ul.y - q.ur.y) <= _TOL and abs(q.ul.x - q.ll.x) <= _TOL
    if len(itens) in (4, 5) and all(i[0] == 'l' for i in itens):
        if not all(abs(i[1][0] - i[2][0]) <= _TOL or abs(i[1][1] - i[2][1]) <= _TOL for i in itens):
            return False
        return _perto(itens[-1][2], itens[0][1])
    return False


def _e_elipse(itens, caixa: fitz.Rect) -> bool:
    if len(itens) < 4 or any(i[0] != 'c' for i in itens) or caixa.width < 1 or caixa.height < 1:
        return False
    if not _perto(itens[-1][4], itens[0][1]):
        return False
    cx, cy = (caixa.x0 + caixa.x1) / 2, (caixa.y0 + caixa.y1) / 2
    a, b = caixa.width / 2, caixa.height / 2
    for i in itens:
        for t in (0.0, 0.5):
            x, y = _ponto_da_curva(i[1], i[2], i[3], i[4], t)
            if abs(((x - cx) / a) ** 2 + ((y - cy) / b) ** 2 - 1) > 0.04:
                return False
    return True


def _raio_do_canto(itens, caixa: fitz.Rect) -> Optional[float]:
    """Raio, se o contorno e um retangulo de cantos redondos (mesmo raio nos quatro)."""
    if len(itens) < 5 or any(i[0] not in ('l', 'c') for i in itens):
        return None
    menor = min(caixa.width, caixa.height)
    if menor < 2:
        return None
    cx, cy = (caixa.x0 + caixa.x1) / 2, (caixa.y0 + caixa.y1) / 2
    cantos: Dict[Tuple[bool, bool], fitz.Rect] = {}
    for i in itens:
        if i[0] != 'c':
            continue
        meio = ((i[1][0] + i[4][0]) / 2, (i[1][1] + i[4][1]) / 2)
        chave = (meio[0] < cx, meio[1] < cy)
        trecho = fitz.Rect(min(p[0] for p in i[1:]), min(p[1] for p in i[1:]),
                           max(p[0] for p in i[1:]), max(p[1] for p in i[1:]))
        cantos[chave] = cantos[chave] | trecho if chave in cantos else trecho
    if len(cantos) != 4:
        return None
    raios = []
    for (esquerda, cima), g in cantos.items():
        x_ok = abs(g.x0 - caixa.x0) <= _TOL if esquerda else abs(g.x1 - caixa.x1) <= _TOL
        y_ok = abs(g.y0 - caixa.y0) <= _TOL if cima else abs(g.y1 - caixa.y1) <= _TOL
        if not (x_ok and y_ok) or abs(g.width - g.height) > _TOL + 0.02 * max(g.width, g.height):
            return None
        raios.append((g.width + g.height) / 2)
    raio = sum(raios) / 4
    if max(raios) - min(raios) > _TOL or raio > menor / 2 + _TOL:
        return None
    for i in itens:
        if i[0] == 'l':
            for p in (i[1], i[2]):
                if min(abs(p[0] - caixa.x0), abs(p[0] - caixa.x1),
                       abs(p[1] - caixa.y0), abs(p[1] - caixa.y1)) > _TOL:
                    return None
    return raio


# ── Escrita no slide ─────────────────────────────────────────────────────────

@dataclass
class Posicao:
    """Leva coordenadas do PDF (pt) para o slide: escala e deslocamento do encaixe da pagina."""
    escala: float
    dx: float
    dy: float

    def x(self, v: float) -> int:
        return int(round((self.dx + v * self.escala) * _EMU_POR_PT))

    def y(self, v: float) -> int:
        return int(round((self.dy + v * self.escala) * _EMU_POR_PT))

    def comprimento(self, v: float) -> int:
        return int(round(v * self.escala * _EMU_POR_PT))


def _hex(cor) -> str:
    return '%02X%02X%02X' % tuple(max(0, min(255, round(c * 255))) for c in cor)


def _preenchimento(cor, opacidade) -> str:
    alfa = ''
    if opacidade is not None and opacidade < 0.999:
        alfa = f'<a:alpha val="{max(0, int(opacidade * 100000))}"/>'
    return f'<a:solidFill><a:srgbClr val="{_hex(cor)}">{alfa}</a:srgbClr></a:solidFill>'


def _tracejado(desenho: dict, largura_pt: float) -> str:
    texto = desenho.get('dashes') or ''
    numeros = [float(n) for n in re.findall(r'[\d.]+', texto.split(']')[0])] if '[' in texto else []
    if not numeros or largura_pt <= 0:
        return ''
    if len(numeros) % 2:
        numeros = numeros * 2
    pares = ''.join(
        f'<a:ds d="{max(1, int(a / largura_pt * 100000))}" sp="{max(1, int(b / largura_pt * 100000))}"/>'
        for a, b in zip(numeros[::2], numeros[1::2]))
    return f'<a:custDash>{pares}</a:custDash>'


def _contorno(desenho: dict, pos: Posicao) -> str:
    if 's' not in desenho['type'] or desenho.get('color') is None:
        return '<a:ln><a:noFill/></a:ln>'
    largura_pt = max(0.25, float(desenho.get('width') or 0.25))
    cap = desenho.get('lineCap') or 0
    cap = cap[0] if isinstance(cap, (tuple, list)) else cap
    juncao = desenho.get('lineJoin') or 0
    juncao_xml = {0: '<a:miter lim="800000"/>', 1: '<a:round/>', 2: '<a:bevel/>'}.get(juncao, '')
    return (f'<a:ln w="{max(1, pos.comprimento(largura_pt))}" cap="{_CAPS.get(cap, "flat")}">'
            f'{_preenchimento(desenho["color"], desenho.get("stroke_opacity"))}'
            f'{_tracejado(desenho, largura_pt)}{juncao_xml}</a:ln>')


def _pinta(elemento, desenho: dict, pos: Posicao) -> None:
    """Troca preenchimento, contorno e efeitos da forma pelos do PDF e tira o estilo do tema
    (que, sem isso, poe sombra e contorno que o PDF nao tem)."""
    estilo = elemento.find(qn('p:style'))
    if estilo is not None:
        elemento.remove(estilo)
    sp_pr = elemento.spPr
    for filho in list(sp_pr):
        if filho.tag not in (qn('a:xfrm'), qn('a:prstGeom'), qn('a:custGeom')):
            sp_pr.remove(filho)
    preenche = 'f' in desenho['type'] and desenho.get('fill') is not None
    fundo = (_preenchimento(desenho['fill'], desenho.get('fill_opacity')) if preenche else '<a:noFill/>')
    for xml in (fundo, _contorno(desenho, pos), '<a:effectLst/>'):
        # declara o namespace "a" na primeira tag para o parse_xml aceitar o trecho solto
        sp_pr.append(parse_xml(re.sub(r'^<(a:\w+)', rf'<\1 {nsdecls("a")}', xml, count=1)))


def _geometria_livre(caminhos, caixa: fitz.Rect, pos: Posicao, preenche: bool, traca: bool) -> str:
    largura = max(1, pos.comprimento(caixa.width))
    altura = max(1, pos.comprimento(caixa.height))

    def pt(p) -> str:
        return (f'<a:pt x="{pos.comprimento(p[0] - caixa.x0)}" '
                f'y="{pos.comprimento(p[1] - caixa.y0)}"/>')

    corpo = ''
    for segmentos, fechado in caminhos:
        for seg in segmentos:
            if seg[0] == 'M':
                corpo += f'<a:moveTo>{pt(seg[1])}</a:moveTo>'
            elif seg[0] == 'L':
                corpo += f'<a:lnTo>{pt(seg[1])}</a:lnTo>'
            else:
                corpo += f'<a:cubicBezTo>{pt(seg[1])}{pt(seg[2])}{pt(seg[3])}</a:cubicBezTo>'
        if fechado or preenche:
            corpo += '<a:close/>'
    atributos = '' if preenche else ' fill="none"'
    atributos += '' if traca else ' stroke="0"'
    return (f'<a:custGeom {nsdecls("a")}><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
            f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
            f'<a:path w="{largura}" h="{altura}"{atributos}>{corpo}</a:path></a:pathLst></a:custGeom>')


def _caixa_do_desenho(caminhos, desenho: dict) -> fitz.Rect:
    pontos = _pontos(caminhos)
    caixa = fitz.Rect(min(p[0] for p in pontos), min(p[1] for p in pontos),
                      max(p[0] for p in pontos), max(p[1] for p in pontos))
    return caixa | fitz.Rect(desenho['rect']) if not fitz.Rect(desenho['rect']).is_empty else caixa


def adicionar_forma(slide, desenho: dict, pos: Posicao, numero: int):
    """Cria no slide a forma que corresponde ao desenho do PDF."""
    caminhos = _caminhos(desenho)
    preenche = 'f' in desenho['type'] and desenho.get('fill') is not None
    traca = 's' in desenho['type'] and desenho.get('color') is not None
    itens = desenho['items']
    caixa = _caixa_do_desenho(caminhos, desenho)

    # traco reto simples: vira uma linha de verdade
    if (not preenche and traca and len(caminhos) == 1 and len(caminhos[0][0]) == 2
            and caminhos[0][0][1][0] == 'L'):
        (_, a), (_, b) = caminhos[0][0]
        linha = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pos.x(a[0]), pos.y(a[1]),
                                           pos.x(b[0]), pos.y(b[1]))
        _pinta(linha._element, desenho, pos)
        linha.name = f'Linha {numero}'
        return linha

    esquerda, topo = pos.x(caixa.x0), pos.y(caixa.y0)
    largura, altura = max(1, pos.comprimento(caixa.width)), max(1, pos.comprimento(caixa.height))
    unico = len(caminhos) == 1 and caminhos[0][1]
    raio = None
    if unico and _e_retangulo(itens, caixa):
        tipo, nome = MSO_SHAPE.RECTANGLE, 'Retângulo'
    elif unico and _e_elipse(itens, caixa):
        tipo, nome = MSO_SHAPE.OVAL, 'Elipse'
    elif unico and (raio := _raio_do_canto(itens, caixa)) is not None:
        tipo, nome = MSO_SHAPE.ROUNDED_RECTANGLE, 'Retângulo arredondado'
    else:
        tipo, nome = MSO_SHAPE.RECTANGLE, 'Forma livre'
    forma = slide.shapes.add_shape(tipo, esquerda, topo, largura, altura)
    if raio is not None:
        forma.adjustments[0] = min(0.5, raio / min(caixa.width, caixa.height))
    if nome == 'Forma livre':
        sp_pr = forma._element.spPr
        sp_pr.replace(sp_pr.find(qn('a:prstGeom')),
                      parse_xml(_geometria_livre(caminhos, caixa, pos, preenche, traca)))
    _pinta(forma._element, desenho, pos)
    forma.name = f'{nome} {numero}'
    return forma


def adicionar_figura(slide, figura: Figura, pos: Posicao, numero: int):
    caixa = figura.caixa
    imagem = slide.shapes.add_picture(BytesIO(figura.dados), pos.x(caixa.x0), pos.y(caixa.y0),
                                      max(1, pos.comprimento(caixa.width)),
                                      max(1, pos.comprimento(caixa.height)))
    imagem.name = f'Imagem {numero}'
    return imagem
